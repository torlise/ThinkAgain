from __future__ import annotations

import base64
import html
import json
import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = ROOT.parent / "逆思維_第1至7章重點整理_更豐富版.docx"
ASSET_DIR = ROOT / "assets"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


@dataclass
class Figure:
    rel_id: str
    src: str
    alt: str


def slugify(text: str, fallback: str) -> str:
    text = text.strip().lower()
    if m := re.search(r"第\s*(\d+)\s*章", text):
        return f"chapter-{m.group(1)}"
    if "一頁掌握" in text:
        return "overview"
    if "整體邏輯" in text:
        return "practice"
    ascii_slug = re.sub(r"[^a-z0-9]+", "-", text).strip("-")
    return ascii_slug or fallback


def clean_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def paragraph_images(paragraph: Paragraph) -> list[str]:
    ids: list[str] = []
    for blip in paragraph._element.xpath(".//a:blip"):
        rel_id = blip.get(f"{{{REL_NS}}}embed")
        if rel_id:
            ids.append(rel_id)
    return ids


def table_images(table: Table) -> list[str]:
    ids: list[str] = []
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                ids.extend(paragraph_images(paragraph))
    return ids


def make_svg_asset(blob: bytes, index: int) -> tuple[str, str]:
    image = Image.open(BytesIO(blob))
    if image.mode not in ("RGB", "L"):
        background = Image.new("RGB", image.size, "white")
        if image.mode == "RGBA":
            background.paste(image, mask=image.getchannel("A"))
        else:
            background.paste(image.convert("RGB"))
        image = background
    else:
        image = image.convert("RGB")

    max_width = 1180
    if image.width > max_width:
        scale = max_width / image.width
        image = image.resize((max_width, int(image.height * scale)), Image.Resampling.LANCZOS)

    quality = 70
    while True:
        out = BytesIO()
        image.save(out, format="JPEG", quality=quality, optimize=True, progressive=True)
        data = base64.b64encode(out.getvalue()).decode("ascii")
        if len(data) < 780_000 or quality <= 44 or image.width <= 900:
            break
        quality -= 8
        if quality <= 54 and image.width > 980:
            scale = 980 / image.width
            image = image.resize((980, int(image.height * scale)), Image.Resampling.LANCZOS)

    filename = f"figure-{index:02d}.svg"
    svg = (
        f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {image.width} {image.height}" '
        f'width="{image.width}" height="{image.height}" role="img">\n'
        f'  <image href="data:image/jpeg;base64,{data}" width="{image.width}" height="{image.height}"/>\n'
        "</svg>\n"
    )
    (ASSET_DIR / filename).write_text(svg, encoding="utf-8")
    return filename, f"{image.width}x{image.height}, q{quality}"


def table_cell_text(cell) -> str:
    return clean_text(" ".join(p.text for p in cell.paragraphs))


def html_table(rows: list[list[str]], class_name: str = "data-table") -> str:
    if not rows:
        return ""
    head, body = rows[0], rows[1:]
    col_count = max(len(r) for r in rows)
    out = [f'<div class="table-wrap"><table class="{class_name}">']
    out.append("<thead><tr>")
    for cell in head:
        out.append(f"<th>{html.escape(cell)}</th>")
    out.append("</tr></thead><tbody>")
    for row in body:
        out.append("<tr>")
        for cell in row + [""] * (col_count - len(row)):
            out.append(f"<td>{html.escape(cell)}</td>")
        out.append("</tr>")
    out.append("</tbody></table></div>")
    return "".join(out)


def callout(text: str, kind: str = "note") -> str:
    text = clean_text(text)
    label = ""
    body = text
    if "｜" in text:
        label, body = text.split("｜", 1)
    else:
        for prefix in ("工作上的實際應用", "記住一句話", "可以養成的習慣", "全書前七章的主軸"):
            if text.startswith(prefix):
                label, body = prefix, text[len(prefix) :].strip()
                break
    label_html = f"<strong>{html.escape(label)}</strong>" if label else ""
    return f'<aside class="callout {kind}">{label_html}<p>{html.escape(body)}</p></aside>'


def flush_bullets(parts: list[str], bullets: list[str]) -> None:
    if bullets:
        parts.append('<ul class="point-list">')
        parts.extend(f"<li>{html.escape(item)}</li>" for item in bullets)
        parts.append("</ul>")
        bullets.clear()


def render_figure(figures: dict[str, Figure], rel_id: str, caption: str = "") -> str:
    fig = figures[rel_id]
    alt = caption or fig.alt
    caption_html = f"<figcaption>{html.escape(caption)}</figcaption>" if caption else ""
    return (
        '<figure class="doc-figure">'
        f'<img src="{html.escape(fig.src)}" alt="{html.escape(alt)}" loading="lazy">'
        f"{caption_html}</figure>"
    )


def render_table(table: Table, figures: dict[str, Figure]) -> str:
    rows = [[table_cell_text(cell) for cell in row.cells] for row in table.rows]
    image_ids = table_images(table)
    row_count = len(rows)
    col_count = len(rows[0]) if rows else 0

    if row_count == 1 and col_count == 1:
        text = rows[0][0]
        kind = "reflection" if text.startswith("反思題") else "note"
        return callout(text, kind)

    if col_count == 2 and image_ids:
        cards = ['<div class="figure-grid">']
        seen: set[str] = set()
        for row in table.rows:
            row_ids = []
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    row_ids.extend(paragraph_images(paragraph))
            caption = clean_text(" ".join(table_cell_text(cell) for cell in row.cells))
            for rel_id in row_ids:
                if rel_id in figures and rel_id not in seen:
                    seen.add(rel_id)
                    cards.append(render_figure(figures, rel_id, caption))
        cards.append("</div>")
        return "".join(cards)

    if col_count == 2 and rows and rows[0] == ["常見思維", "重新思考後"]:
        cards = ['<div class="compare-grid" aria-label="從舊思維到重新思考">']
        for before, after in rows[1:]:
            cards.append(
                '<article class="compare-card">'
                f'<p class="before">{html.escape(before)}</p>'
                f'<p class="after">{html.escape(after)}</p>'
                "</article>"
            )
        cards.append("</div>")
        return "".join(cards)

    return html_table(rows, "summary-table" if col_count == 3 else "data-table")


def build_site() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document(DOCX_PATH)

    ordered_image_ids: list[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            ordered_image_ids.extend(paragraph_images(Paragraph(child, doc)))
        elif child.tag == qn("w:tbl"):
            ordered_image_ids.extend(table_images(Table(child, doc)))

    figures: dict[str, Figure] = {}
    image_index = 1
    asset_meta = []
    for rel_id in ordered_image_ids:
        if rel_id in figures:
            continue
        part = doc.part.related_parts[rel_id]
        filename, meta = make_svg_asset(part.blob, image_index)
        figures[rel_id] = Figure(rel_id=rel_id, src=f"assets/{filename}", alt=f"《逆思維》圖解 {image_index}")
        asset_meta.append({"file": filename, "source": str(part.partname), "meta": meta})
        image_index += 1

    parts: list[str] = []
    nav: list[tuple[str, str]] = []
    bullets: list[str] = []
    pending_label = ""
    open_section = False
    section_index = 0

    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, doc)
            text = clean_text(paragraph.text)
            style = paragraph.style.name if paragraph.style else ""
            img_ids = paragraph_images(paragraph)

            if not text and not img_ids:
                continue

            if style == "List Bullet":
                bullets.append(text)
                continue

            flush_bullets(parts, bullets)

            if style == "Section Label":
                pending_label = text
                continue

            if style == "Heading 1":
                if open_section:
                    parts.append("</section>")
                section_index += 1
                section_id = slugify(text, f"section-{section_index}")
                nav.append((section_id, text))
                section_class = "chapter-section overview" if section_id == "overview" else "chapter-section"
                parts.append(f'<section id="{section_id}" class="{section_class}" data-title="{html.escape(text)}">')
                if pending_label:
                    parts.append(f'<p class="section-kicker">{html.escape(pending_label)}</p>')
                    pending_label = ""
                parts.append(f"<h2>{html.escape(text)}</h2>")
                open_section = True
                continue

            if not open_section:
                if text:
                    parts.append(f'<p class="eyebrow">{html.escape(text)}</p>' if "Think Again" in text else f"<p>{html.escape(text)}</p>")
                continue

            if style == "Heading 2":
                parts.append(f"<h3>{html.escape(text)}</h3>")
            else:
                if text:
                    parts.append(f"<p>{html.escape(text)}</p>")

            for rel_id in img_ids:
                if rel_id in figures:
                    parts.append(render_figure(figures, rel_id))

        elif child.tag == qn("w:tbl"):
            flush_bullets(parts, bullets)
            parts.append(render_table(Table(child, doc), figures))

    flush_bullets(parts, bullets)
    if open_section:
        parts.append("</section>")

    nav_html = "\n".join(
        f'<a href="#{html.escape(section_id)}">{html.escape(title)}</a>' for section_id, title in nav
    )
    generated_content = "\n".join(parts)

    index = f"""<!doctype html>
<html lang="zh-Hant">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>逆思維 Think Again｜第 1～7 章讀書會閱讀器</title>
  <meta name="description" content="《逆思維 Think Again》第 1～7 章重點整理的互動式讀書會閱讀器。">
  <link rel="stylesheet" href="styles.css">
</head>
<body>
  <div class="progress" aria-hidden="true"><span></span></div>
  <header class="site-header">
    <div>
      <p class="eyebrow">Adam Grant / Think Again</p>
      <h1>《逆思維 Think Again》</h1>
      <p class="subtitle">第 1～7 章重點整理，為讀書會重新排成可瀏覽、可討論、可複習的網頁閱讀方式。</p>
    </div>
    <a class="skip-link" href="#overview">開始閱讀</a>
  </header>

  <div class="reader-shell">
    <nav class="chapter-nav" aria-label="章節導覽">
      <div class="nav-title">章節</div>
      {nav_html}
    </nav>
    <main class="reader-content">
      {generated_content}
    </main>
  </div>

  <button class="to-top" type="button" aria-label="回到頁首">↑</button>
  <script src="app.js"></script>
</body>
</html>
"""
    (ROOT / "index.html").write_text(index, encoding="utf-8")

    metadata = {
        "source_docx": DOCX_PATH.name,
        "figures": asset_meta,
        "sections": [{"id": section_id, "title": title} for section_id, title in nav],
    }
    (ROOT / "assets" / "manifest.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8"
    )


if __name__ == "__main__":
    build_site()
